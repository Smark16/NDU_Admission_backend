"""Excel export for test/exam registration sheets (openpyxl) -- same eligible
student list as the PDF sheet, for staff who want to print/edit in Excel."""
from __future__ import annotations

import io


def render_test_exam_registration_excel(
    *,
    kind: str,
    course_code: str,
    course_name: str,
    programme_name: str,
    semester_name: str,
    min_pct: float,
    students: list[dict],
) -> tuple[bytes, str]:
    """Return (xlsx_bytes, filename)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from django.utils import timezone

    kind_label = "Test" if kind == "test" else "Exam"

    wb = Workbook()
    ws = wb.active
    ws.title = f"{kind_label} Registration"

    title_font = Font(bold=True, size=14, color="1F3A5F")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="3E397B")

    ws["A1"] = f"{kind_label} Registration Sheet"
    ws["A1"].font = title_font
    ws.merge_cells("A1:F1")

    info_rows = [
        ("Course", f"{course_code} — {course_name}"),
        ("Programme", programme_name or "—"),
        ("Semester", semester_name or "—"),
        ("Minimum tuition paid", f"{min_pct:.0f}% (or active temporary access pass)"),
        ("Generated", timezone.localtime().strftime("%d %B %Y %I:%M %p")),
    ]
    row = 3
    for label, value in info_rows:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1

    header_row = row + 1
    headers = ["#", "Reg No", "Student Name", "% Paid", "Basis", "Signature"]
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    r = header_row + 1
    for i, s in enumerate(students, start=1):
        ws.cell(row=r, column=1, value=i)
        ws.cell(row=r, column=2, value=s.get("reg_no", ""))
        ws.cell(row=r, column=3, value=s.get("name", ""))
        ws.cell(row=r, column=4, value=s.get("percentage_paid", ""))
        ws.cell(row=r, column=5, value=s.get("basis", ""))
        ws.cell(row=r, column=6, value="")
        r += 1

    if not students:
        ws.cell(row=r, column=1, value="No eligible students.")
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Invigilator / Lecturer signature:").font = Font(bold=True)
    r += 1
    ws.cell(row=r, column=1, value="Date:").font = Font(bold=True)

    widths = {"A": 6, "B": 18, "C": 32, "D": 10, "E": 16, "F": 24}
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    code = "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in (course_code or "course"))
    date_part = timezone.localdate().strftime("%Y%m%d")
    filename = f"{kind_label}Registration_{code.strip('_') or 'course'}_{date_part}.xlsx"

    return buffer.read(), filename

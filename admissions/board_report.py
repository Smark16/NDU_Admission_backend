"""Board admissions/registration report -- Tables 1-4, computed live per
intake batch, matching the Academic Registrar's approved report structure.

Table 1: Summary by campus.
Table 2: Applicants/Admitted/Registered by Faculty/School (HEC carved out).
Table 3: HEC students per campus, by subject area.
Table 4: Statistics by Faculty per Programme and campus (delivery-mode
         variants -- Day/Weekend/Main/etc -- consolidated into one row per
         campus).
"""
from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import AdmittedStudent, Application, ApplicationProgramChoice, Batch

CAMPUS_LABELS = {
    "NDEJJE UNIVERSITY LUWERO CAMPUS": "Main Campus",
    "NDEJJE UNIVERSITY KAMPALA CAMPUS": "Kampala Campus",
}
FACULTY_LABELS = {
    "Faculty of Engineering & Survey": "Faculty of Engineering",
    "Faculty of Business Administration & Management": "Faculty of Business Admn & Mgt",
    "Faculty of Education & Humanities": "Faculty of Education & Humanities",
    "Faculty of Science & Computing": "Faculty of Science and computing",
    "Faculty of Environment & Agriculture": "Environment and Agricultural Sciences",
    "Faculty of Health Sciences": "Faculty of Health",
    "Faculty of Law": "Faculty of Law",
    "Graduate School": "School of post Graduate",
}
TABLE2_ORDER = [
    "Faculty of Engineering",
    "Faculty of Business Admn & Mgt",
    "Faculty of Education & Humanities",
    "Faculty of Science and computing",
    "Environment and Agricultural Sciences",
    "Faculty of Health",
    "Faculty of Law",
    "School of post Graduate",
    "HEC Students",
]

MODE_SUFFIX_RE = re.compile(
    r"\s*[-–]?\s*\(?\b(day|weekend|wkd|main|inservice|in-service|international|"
    r"distance learning)\b\)?\s*$",
    re.IGNORECASE,
)

# Known duplicate / near-duplicate programme names (data-entry variants of
# the same real programme) folded together for Table 4.
PROGRAMME_ALIASES = {
    "Bachelor of Forest Science and Environmental Management": "Bachelor of Forest Science and Environment Management",
    "Postgraduate Diploma in Education (PDE)": "Postgraduate Diploma in Education",
}


def _clean_programme_name(name: str) -> str:
    n = (name or "").strip()
    while True:
        new_n = MODE_SUFFIX_RE.sub("", n).strip(" -–")
        if new_n == n:
            break
        n = new_n
    n = re.sub(r"\s+", " ", n)
    return PROGRAMME_ALIASES.get(n, n)


def _is_hec(name: str) -> bool:
    return "higher education certificate" in (name or "").lower()


def _campus_label(campus) -> str:
    if campus is None:
        return "(no campus)"
    return CAMPUS_LABELS.get(campus.name, campus.name)


def _faculty_label(faculty) -> str:
    if faculty is None:
        return "(no faculty)"
    return FACULTY_LABELS.get(faculty.name, faculty.name)


def _rate(numerator: int, denominator: int) -> float:
    return round((numerator / denominator * 100), 1) if denominator else 0.0


def _table1(batch_id: int) -> dict:
    applicants = {}
    for row in Application.objects.filter(batch_id=batch_id).values("campus__name"):
        label = CAMPUS_LABELS.get(row["campus__name"], row["campus__name"] or "(no campus)")
        applicants[label] = applicants.get(label, 0) + 1

    admitted = {}
    registered = {}
    course_registered = {}
    for s in AdmittedStudent.objects.filter(admitted_batch_id=batch_id).select_related("admitted_campus"):
        label = _campus_label(s.admitted_campus)
        admitted[label] = admitted.get(label, 0) + 1
        if s.accounts_registration_cleared:
            registered[label] = registered.get(label, 0) + 1
        if s.is_registered:
            course_registered[label] = course_registered.get(label, 0) + 1

    rows = []
    tot_a = tot_ad = tot_r = tot_cr = 0
    for label in ["Main Campus", "Kampala Campus"]:
        a = applicants.get(label, 0)
        ad = admitted.get(label, 0)
        r = registered.get(label, 0)
        cr = course_registered.get(label, 0)
        tot_a += a
        tot_ad += ad
        tot_r += r
        tot_cr += cr
        rows.append({
            "campus": label, "applicants": a, "admitted": ad,
            "registered": r, "rate": _rate(r, ad),
            "course_registered": cr, "course_registered_rate": _rate(cr, ad),
        })
    total = {
        "campus": "Total", "applicants": tot_a, "admitted": tot_ad,
        "registered": tot_r, "rate": _rate(tot_r, tot_ad),
        "course_registered": tot_cr, "course_registered_rate": _rate(tot_cr, tot_ad),
    }
    return {"rows": rows, "total": total}


def _table2(batch_id: int) -> dict:
    applicants = {}
    apps_with_choice_ids = set()
    for c in (
        ApplicationProgramChoice.objects.filter(application__batch_id=batch_id, choice_order=1)
        .select_related("program", "program__faculty")
    ):
        apps_with_choice_ids.add(c.application_id)
        key = "HEC Students" if _is_hec(c.program.name) else _faculty_label(c.program.faculty)
        applicants[key] = applicants.get(key, 0) + 1

    no_choice_apps = (
        Application.objects.filter(batch_id=batch_id)
        .exclude(id__in=apps_with_choice_ids)
        .select_related("admission", "admission__admitted_program", "admission__admitted_program__faculty")
    )
    no_program_at_all = 0
    for app in no_choice_apps:
        admission = getattr(app, "admission", None)
        if admission is None or admission.admitted_program_id is None:
            no_program_at_all += 1
            continue
        program = admission.admitted_program
        key = "HEC Students" if _is_hec(program.name) else _faculty_label(program.faculty)
        applicants[key] = applicants.get(key, 0) + 1

    admitted = {}
    registered = {}
    course_registered = {}
    for s in AdmittedStudent.objects.filter(admitted_batch_id=batch_id).select_related("admitted_program", "admitted_program__faculty"):
        key = "HEC Students" if _is_hec(s.admitted_program.name) else _faculty_label(s.admitted_program.faculty)
        admitted[key] = admitted.get(key, 0) + 1
        if s.accounts_registration_cleared:
            registered[key] = registered.get(key, 0) + 1
        if s.is_registered:
            course_registered[key] = course_registered.get(key, 0) + 1

    rows = []
    tot_a = tot_ad = tot_r = tot_cr = 0
    for label in TABLE2_ORDER:
        a = applicants.get(label, 0)
        ad = admitted.get(label, 0)
        r = registered.get(label, 0)
        cr = course_registered.get(label, 0)
        tot_a += a
        tot_ad += ad
        tot_r += r
        tot_cr += cr
        rows.append({
            "faculty": label, "applicants": a, "admitted": ad,
            "registered": r, "rate": _rate(r, ad),
            "course_registered": cr, "course_registered_rate": _rate(cr, ad),
        })
    if no_program_at_all:
        rows.append({
            "faculty": "(no programme on record)", "applicants": no_program_at_all, "admitted": 0,
            "registered": 0, "rate": 0, "course_registered": 0, "course_registered_rate": 0,
        })
        tot_a += no_program_at_all
    total = {
        "faculty": "Total", "applicants": tot_a, "admitted": tot_ad,
        "registered": tot_r, "rate": _rate(tot_r, tot_ad),
        "course_registered": tot_cr, "course_registered_rate": _rate(tot_cr, tot_ad),
    }
    return {"rows": rows, "total": total}


def _table3_classify(program_name: str) -> str:
    name = program_name.lower()
    if "international" in name:
        return "International"
    if "business" in name:
        return "Business"
    if "physical sciences" in name:
        return "Physical sciences"
    if "biological sciences" in name:
        return "Biological sciences"
    if "humanities" in name:
        return "Humanities"
    return "(unclassified)"


def _table3(batch_id: int) -> dict:
    counts = {}
    for s in (
        AdmittedStudent.objects.filter(admitted_batch_id=batch_id, admitted_program__name__icontains="Higher Education Certificate")
        .select_related("admitted_program", "admitted_campus")
    ):
        subject = _table3_classify(s.admitted_program.name)
        if subject == "International":
            key = ("International", "International")
        else:
            key = (_campus_label(s.admitted_campus), subject)
        counts[key] = counts.get(key, 0) + 1

    def _rows_for(campus_label: str, order: list[str]) -> tuple[list[dict], int]:
        out = []
        total = 0
        for subj in order:
            n = counts.pop((campus_label, subj), 0)
            total += n
            out.append({"subject": subj, "count": n})
        return out, total

    kla_rows, kla_total = _rows_for("Kampala Campus", ["Biological sciences", "Business", "Humanities", "Physical sciences"])
    main_rows, main_total = _rows_for("Main Campus", ["Biological sciences", "Physical sciences", "Business"])
    international = counts.pop(("International", "International"), 0)

    unplaced = [{"campus": c, "subject": s, "count": n} for (c, s), n in counts.items()]

    return {
        "kampala": {"rows": kla_rows, "total": kla_total},
        "main": {"rows": main_rows, "total": main_total},
        "international": international,
        "grand_total": kla_total + main_total + international,
        "unplaced": unplaced,
    }


def _table4(batch_id: int) -> dict:
    admitted_counts = {}
    registered_counts = {}
    course_registered_counts = {}
    for s in (
        AdmittedStudent.objects.filter(admitted_batch_id=batch_id)
        .exclude(admitted_program__name__icontains="Higher Education Certificate")
        .select_related("admitted_program", "admitted_program__faculty", "admitted_campus")
    ):
        campus = _campus_label(s.admitted_campus)
        fac = _faculty_label(s.admitted_program.faculty)
        prog = _clean_programme_name(s.admitted_program.name)
        key = (campus, fac, prog)
        admitted_counts[key] = admitted_counts.get(key, 0) + 1
        if s.accounts_registration_cleared:
            registered_counts[key] = registered_counts.get(key, 0) + 1
        if s.is_registered:
            course_registered_counts[key] = course_registered_counts.get(key, 0) + 1

    applicant_counts = {}
    apps_with_choice_ids = set()
    for c in (
        ApplicationProgramChoice.objects.filter(application__batch_id=batch_id, choice_order=1)
        .select_related("program", "program__faculty", "application", "application__campus")
    ):
        apps_with_choice_ids.add(c.application_id)
        if _is_hec(c.program.name):
            continue
        campus = _campus_label(c.application.campus)
        fac = _faculty_label(c.program.faculty)
        prog = _clean_programme_name(c.program.name)
        key = (campus, fac, prog)
        applicant_counts[key] = applicant_counts.get(key, 0) + 1

    no_choice_apps = (
        Application.objects.filter(batch_id=batch_id)
        .exclude(id__in=apps_with_choice_ids)
        .select_related("campus", "admission", "admission__admitted_program", "admission__admitted_program__faculty")
    )
    for app in no_choice_apps:
        admission = getattr(app, "admission", None)
        if admission is None or admission.admitted_program_id is None:
            continue
        program = admission.admitted_program
        if _is_hec(program.name):
            continue
        campus = _campus_label(app.campus)
        fac = _faculty_label(program.faculty)
        prog = _clean_programme_name(program.name)
        key = (campus, fac, prog)
        applicant_counts[key] = applicant_counts.get(key, 0) + 1

    all_keys = set(admitted_counts) | set(applicant_counts) | set(registered_counts) | set(course_registered_counts)
    rows = []
    for campus, fac, prog in sorted(all_keys):
        a = applicant_counts.get((campus, fac, prog), 0)
        ad = admitted_counts.get((campus, fac, prog), 0)
        r = registered_counts.get((campus, fac, prog), 0)
        cr = course_registered_counts.get((campus, fac, prog), 0)
        rows.append({
            "campus": campus, "faculty": fac, "programme": prog,
            "applicants": a, "admitted": ad, "registered": r, "course_registered": cr,
        })

    total = {
        "applicants": sum(applicant_counts.values()),
        "admitted": sum(admitted_counts.values()),
        "registered": sum(registered_counts.values()),
        "course_registered": sum(course_registered_counts.values()),
    }
    return {"rows": rows, "total": total}


def build_board_report(batch_id: int) -> dict:
    batch = Batch.objects.get(pk=batch_id)
    return {
        "batch": {"id": batch.id, "name": batch.name, "academic_year": batch.academic_year},
        "table1": _table1(batch_id),
        "table2": _table2(batch_id),
        "table3": _table3(batch_id),
        "table4": _table4(batch_id),
    }


# --- Excel export -----------------------------------------------------------

_HEADER_FILL = PatternFill("solid", fgColor="3E397B")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, size=13, color="2D2960")
_TOTAL_FONT = Font(bold=True)
_THIN = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT = Alignment(horizontal="left", vertical="center")


def _write_header_row(sheet, row: int, headers: list[str]):
    for col, label in enumerate(headers, start=1):
        cell = sheet.cell(row, col, label)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER
        cell.border = _THIN


def _write_data_row(sheet, row: int, values: list, *, bold: bool = False):
    for col, value in enumerate(values, start=1):
        cell = sheet.cell(row, col, value)
        cell.border = _THIN
        cell.alignment = _LEFT if col <= 2 else _CENTER
        if bold:
            cell.font = _TOTAL_FONT


def _autosize(sheet, n_cols: int):
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        max_len = max(
            (len(str(sheet.cell(r, col).value or "")) for r in range(1, sheet.max_row + 1)),
            default=10,
        )
        sheet.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)


def board_report_xlsx(payload: dict) -> bytes:
    wb = Workbook()
    batch_name = payload["batch"]["name"]

    # Table 1
    ws1 = wb.active
    ws1.title = "Table 1 - By Campus"
    ws1.merge_cells("A1:E1")
    ws1["A1"] = f"Table 1: Summary by Campus — {batch_name}"
    ws1["A1"].font = _TITLE_FONT
    _write_header_row(ws1, 3, ["Campus", "Applicants", "Admitted", "Registered", "Course Registered"])
    row = 4
    for r in payload["table1"]["rows"]:
        _write_data_row(ws1, row, [r["campus"], r["applicants"], r["admitted"], r["registered"], r["course_registered"]])
        row += 1
    t = payload["table1"]["total"]
    _write_data_row(ws1, row, [t["campus"], t["applicants"], t["admitted"], t["registered"], t["course_registered"]], bold=True)
    _autosize(ws1, 5)

    # Table 2
    ws2 = wb.create_sheet("Table 2 - By Faculty")
    ws2.merge_cells("A1:E1")
    ws2["A1"] = f"Table 2: Applicants, Admitted, Registered by Faculty/School — {batch_name}"
    ws2["A1"].font = _TITLE_FONT
    _write_header_row(ws2, 3, ["Faculty/School", "Applicants", "Admitted", "Registered", "Course Registered"])
    row = 4
    for r in payload["table2"]["rows"]:
        _write_data_row(ws2, row, [r["faculty"], r["applicants"], r["admitted"], r["registered"], r["course_registered"]])
        row += 1
    t = payload["table2"]["total"]
    _write_data_row(ws2, row, [t["faculty"], t["applicants"], t["admitted"], t["registered"], t["course_registered"]], bold=True)
    _autosize(ws2, 5)

    # Table 3
    ws3 = wb.create_sheet("Table 3 - HEC Students")
    ws3.merge_cells("A1:B1")
    ws3["A1"] = f"Table 3: Number of HEC Students per campus — {batch_name}"
    ws3["A1"].font = _TITLE_FONT
    row = 3
    _write_header_row(ws3, row, ["Kampala campus", "Total"])
    row += 1
    for r in payload["table3"]["kampala"]["rows"]:
        _write_data_row(ws3, row, [r["subject"], r["count"]])
        row += 1
    _write_data_row(ws3, row, ["Total", payload["table3"]["kampala"]["total"]], bold=True)
    row += 2
    _write_header_row(ws3, row, ["Main campus", "Total"])
    row += 1
    for r in payload["table3"]["main"]["rows"]:
        _write_data_row(ws3, row, [r["subject"], r["count"]])
        row += 1
    _write_data_row(ws3, row, ["Total", payload["table3"]["main"]["total"]], bold=True)
    row += 2
    _write_data_row(ws3, row, ["International Students", payload["table3"]["international"]], bold=True)
    row += 1
    _write_data_row(ws3, row, ["Grand Total (Kampala + Main + International)", payload["table3"]["grand_total"]], bold=True)
    _autosize(ws3, 2)

    # Table 4
    ws4 = wb.create_sheet("Table 4 - By Programme")
    ws4.merge_cells("A1:G1")
    ws4["A1"] = f"Table 4: Statistics by Faculty per Programme and campus — {batch_name}"
    ws4["A1"].font = _TITLE_FONT
    _write_header_row(ws4, 3, ["Campus", "Faculty/School", "Programme", "Applicants", "Admitted", "Registered", "Course Registered"])
    row = 4
    for r in payload["table4"]["rows"]:
        _write_data_row(ws4, row, [
            r["campus"], r["faculty"], r["programme"], r["applicants"], r["admitted"],
            r["registered"], r["course_registered"],
        ])
        row += 1
    t = payload["table4"]["total"]
    _write_data_row(ws4, row, ["", "", "Total", t["applicants"], t["admitted"], t["registered"], t["course_registered"]], bold=True)
    _autosize(ws4, 7)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

"""Commitment fee report JSON + Excel (admitted students with ≥ UGX 150,000 paid)."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from django.http import HttpResponse
from rest_framework.permissions import BasePermission, IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.erp_drf_permissions import user_has_any_erp_perm
from accounts.super_admin import user_is_super_admin
from admissions.commitment_fee_report import (
    build_commitment_fee_report,
    commitment_fee_report_filter_options,
)


class CommitmentFeeReportPermission(BasePermission):
    message = "You do not have permission to view the commitment fee report."

    def has_permission(self, request, view):
        u = request.user
        if not u or not u.is_authenticated:
            return False
        if user_is_super_admin(u):
            return True
        if u.has_perm("admissions.view_admittedstudent"):
            return True
        if u.has_perm("AdmissionReports.view_admissionreports"):
            return True
        return user_has_any_erp_perm(u, "access_reports", "access_finance")


def _parse_int(raw) -> int | None:
    try:
        value = int(str(raw or "").strip())
        return value if value > 0 else None
    except (TypeError, ValueError):
        return None


def _params(request) -> dict:
    sort = (request.query_params.get("program_sort") or "name").strip().lower()
    if sort not in (
        "name",
        "students_asc",
        "students_desc",
        "paid_asc",
        "paid_desc",
    ):
        sort = "name"
    return {
        "program_id": _parse_int(request.query_params.get("program")),
        "campus_id": _parse_int(request.query_params.get("campus")),
        "batch_id": _parse_int(request.query_params.get("batch")),
        "search": (request.query_params.get("search") or "").strip(),
        "min_students": _parse_int(request.query_params.get("min_students")),
        "max_students": _parse_int(request.query_params.get("max_students")),
        "program_sort": sort,
    }


def _ugx(n) -> str:
    if n is None:
        return ""
    return f"{float(n):,.0f}"


class CommitmentFeeReportView(APIView):
    permission_classes = [IsAuthenticated, CommitmentFeeReportPermission]

    def get(self, request):
        params = _params(request)
        try:
            data = build_commitment_fee_report(request.user, params)
            data["filters"] = commitment_fee_report_filter_options(request.user)
        except Exception:
            import logging

            logging.getLogger(__name__).exception("commitment fee report failed")
            return Response(
                {"detail": "Could not build the commitment fee report. Check server logs."},
                status=500,
            )
        data["applied"] = {
            "program": params["program_id"],
            "campus": params["campus_id"],
            "batch": params["batch_id"],
            "search": params["search"] or None,
            "min_students": params["min_students"],
            "max_students": params["max_students"],
            "program_sort": params["program_sort"],
        }
        return Response(data)


class CommitmentFeeReportExcelView(APIView):
    permission_classes = [IsAuthenticated, CommitmentFeeReportPermission]

    def get(self, request):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        params = _params(request)
        sheet = (request.query_params.get("sheet") or "all").strip().lower()
        data = build_commitment_fee_report(request.user, params)
        totals = data.get("totals") or {}
        wb = Workbook()
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1e1b4b", fill_type="solid")
        thin = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )

        def write_sheet(ws, headers, rows):
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", wrap_text=True)
                cell.border = thin
                ws.column_dimensions[get_column_letter(col)].width = max(
                    14, min(36, len(str(header)) + 4)
                )
            for r_i, row in enumerate(rows, 2):
                for c_i, value in enumerate(row, 1):
                    cell = ws.cell(row=r_i, column=c_i, value=value)
                    cell.border = thin
            ws.freeze_panes = "A2"
            if rows:
                ws.auto_filter.ref = ws.dimensions

        programme_rows = [
            [
                row.get("program") or "",
                row.get("program_code") or "",
                row.get("faculty") or "",
                row.get("students_count") or 0,
                _ugx(row.get("total_paid_ugx")),
            ]
            for row in data.get("by_program") or []
        ]

        if sheet in ("programmes", "by_programme", "by_program"):
            ws = wb.active
            ws.title = "By programme"
            write_sheet(
                ws,
                ["Programme", "Code", "Faculty", "Students", "Total paid (UGX)"],
                programme_rows,
            )
            stamp = datetime.now().strftime("%Y-%m-%d")
            filename = f"commitment_fee_programmes_{stamp}.xlsx"
        else:
            summary = wb.active
            summary.title = "Summary"
            write_sheet(
                summary,
                ["Metric", "Value"],
                [
                    ["Students (commitment met)", totals.get("students_count") or 0],
                    ["Programmes", totals.get("programs_count") or 0],
                    ["Total paid (UGX)", _ugx(totals.get("total_paid_ugx"))],
                    ["Commitment threshold (UGX)", _ugx(totals.get("commitment_threshold"))],
                    ["Min students filter", params.get("min_students") or "—"],
                    ["Max students filter", params.get("max_students") or "—"],
                    ["Programme sort", params.get("program_sort") or "name"],
                ],
            )

            by_program = wb.create_sheet("By programme")
            write_sheet(
                by_program,
                ["Programme", "Code", "Faculty", "Students", "Total paid (UGX)"],
                programme_rows,
            )

            students = wb.create_sheet("Students")
            write_sheet(
                students,
                [
                    "Name",
                    "Student ID",
                    "Reg no",
                    "SchoolPay code",
                    "Programme",
                    "Code",
                    "Faculty",
                    "Campus",
                    "Intake",
                    "Academic year",
                    "Paid (UGX)",
                    "Threshold (UGX)",
                    "Flagged admission fee paid",
                    "Admission fee paid at",
                ],
                [
                    [
                        row.get("name") or "",
                        row.get("student_id") or "",
                        row.get("reg_no") or "",
                        row.get("schoolpay_code") or "",
                        row.get("program") or "",
                        row.get("program_code") or "",
                        row.get("faculty") or "",
                        row.get("campus") or "",
                        row.get("intake") or "",
                        row.get("academic_year") or "",
                        _ugx(row.get("paid_ugx")),
                        _ugx(row.get("commitment_threshold")),
                        "Yes" if row.get("admission_fee_paid") else "No",
                        (row.get("admission_fee_paid_at") or "")[:19].replace("T", " "),
                    ]
                    for row in data.get("students") or []
                ],
            )
            stamp = datetime.now().strftime("%Y-%m-%d")
            filename = f"commitment_fee_report_{stamp}.xlsx"

        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        response = HttpResponse(
            bio.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
